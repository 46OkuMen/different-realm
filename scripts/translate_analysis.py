import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('DiffRealm_Text.xlsx')

# Stats storage
sheet_stats = []
grand_total_strings = 0
grand_total_translated = 0
grand_total_untranslated = 0
grand_total_empty = 0

print('='*100)
print('TRANSLATION PROGRESS ANALYSIS')
print('='*100)

# Process each worksheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    total_strings = 0
    translated_count = 0
    untranslated_count = 0
    empty_count = 0
    
    examples_translated = []
    examples_untranslated = []
    
    # Iterate through rows (skip header if present, start from row 2)
    for row in ws.iter_rows(min_row=2, values_only=False):
        col_e = row[4]  # Column E (Japanese)
        col_g = row[6]  # Column G (English)
        
        # Check if column E has text (Japanese)
        jp_text = col_e.value if col_e else None
        en_text = col_g.value if col_g else None
        
        if jp_text and isinstance(jp_text, str) and jp_text.strip():
            total_strings += 1
            
            # Check if translated (G != E and not empty)
            if en_text and isinstance(en_text, str) and en_text.strip():
                if en_text.strip() != jp_text.strip():
                    translated_count += 1
                    if len(examples_translated) < 5:
                        examples_translated.append((jp_text.strip()[:50], en_text.strip()[:50]))
                else:
                    # G equals E (not translated)
                    untranslated_count += 1
                    if len(examples_untranslated) < 5:
                        examples_untranslated.append(jp_text.strip()[:50])
            else:
                # G is empty
                empty_count += 1
                if len(examples_untranslated) < 5:
                    examples_untranslated.append(jp_text.strip()[:50])
    
    # Calculate percentage
    pct_translated = (translated_count / total_strings * 100) if total_strings > 0 else 0
    
    sheet_stats.append({
        'name': sheet_name,
        'total': total_strings,
        'translated': translated_count,
        'untranslated': untranslated_count + empty_count,
        'pct': pct_translated,
        'examples_translated': examples_translated,
        'examples_untranslated': examples_untranslated
    })
    
    grand_total_strings += total_strings
    grand_total_translated += translated_count
    grand_total_untranslated += untranslated_count
    grand_total_empty += empty_count

# Print summary table
print('\n{:<30} {:<15} {:<15} {:<15} {:<10}'.format('Sheet Name', 'Total Strings', 'Translated', 'Untranslated', '% Done'))
print('-' * 90)

for stats in sheet_stats:
    print('{:<30} {:<15} {:<15} {:<15} {:<9.1f}%'.format(
        stats['name'], 
        stats['total'], 
        stats['translated'], 
        stats['untranslated'], 
        stats['pct']
    ))

print('-' * 90)
grand_pct = (grand_total_translated / grand_total_strings * 100) if grand_total_strings > 0 else 0
print('{:<30} {:<15} {:<15} {:<15} {:<9.1f}%'.format(
    'GRAND TOTAL', 
    grand_total_strings, 
    grand_total_translated, 
    grand_total_untranslated + grand_total_empty, 
    grand_pct
))

# Show examples
print('\n')
print('='*100)
print('EXAMPLE TRANSLATIONS (First 5 where EN != JP):')
print('='*100)
for stats in sheet_stats:
    if stats['examples_translated']:
        print('\n[{}]'.format(stats['name']))
        for jp, en in stats['examples_translated']:
            print('  JP: {}'.format(jp))
            print('  EN: {}'.format(en))
            print()

print('\n')
print('='*100)
print('EXAMPLE UNTRANSLATED STRINGS (First 5 per sheet):')
print('='*100)
for stats in sheet_stats:
    if stats['examples_untranslated']:
        print('\n[{}]'.format(stats['name']))
        for jp in stats['examples_untranslated']:
            print('  {}'.format(jp))
