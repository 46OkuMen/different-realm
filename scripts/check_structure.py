import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('DiffRealm_Text.xlsx')

# Check first sheet structure
ws = wb.active
print('First sheet name:', ws.title)
print('Max row:', ws.max_row)
print('Max column:', ws.max_column)
print('\nFirst 5 rows (showing columns A-H):')
print('-' * 120)

for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
    print(f'Row {i+1}:')
    for j, cell in enumerate(row[:8]):
        col_letter = chr(65 + j)  # A=65
        print(f'  {col_letter}: {str(cell)[:60]}')
    print()

print('\nChecking column headers (row 1):')
for i in range(8):
    cell = ws.cell(row=1, column=i+1)
    col_letter = chr(65 + i)
    print(f'{col_letter}1: {cell.value}')
