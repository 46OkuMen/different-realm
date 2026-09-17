#!/usr/bin/env python
"""Merge English translations from a shared Google Sheet into DiffRealm_Text.xlsx.

This preserves the local workbook structure (sheet order, formulas, suffix/comments
columns) and only updates the English column.

Recommended workflow for private Google Sheets:
    1. In Google Sheets, use File -> Download -> Microsoft Excel (.xlsx)
    2. Run:
       python sync_google_sheet.py path\\to\\downloaded_sheet.xlsx

If the sheet is public or link-shared for anonymous access, a Google Sheets URL or
sheet ID can be used directly:
    python sync_google_sheet.py "https://docs.google.com/spreadsheets/d/<ID>/edit"
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook


DEFAULT_TARGET = Path("DiffRealm_Text.xlsx")
GOOGLE_SHEET_ID_RE = re.compile(r"(?:/d/|key=)([A-Za-z0-9_-]+)")


@dataclass
class RowRecord:
    sheet: str
    row_number: int
    offset_key: str
    block_key: str
    japanese_key: str
    japanese_text: str
    english_text: str


@dataclass
class SourceIndex:
    exact_rows: dict
    occurrence_rows: dict
    sheet_names: set
    total_rows: int
    translated_rows: int


def normalize_key(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().replace("\r\n", "\n").replace("\r", "\n")


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def is_translated(japanese, english):
    jp_key = normalize_key(japanese)
    en_key = normalize_key(english)
    return bool(en_key) and en_key != jp_key


def extract_google_sheet_id(source_spec):
    if GOOGLE_SHEET_ID_RE.search(source_spec):
        return GOOGLE_SHEET_ID_RE.search(source_spec).group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{20,}", source_spec):
        return source_spec
    return None


def download_google_sheet_xlsx(source_spec):
    sheet_id = extract_google_sheet_id(source_spec)
    if not sheet_id:
        return None

    export_url = (
        f"https://docs.google.com/spreadsheets/d/{sheet_id}/export"
        f"?format=xlsx"
    )

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_path = Path(tmp.name)
    tmp.close()

    try:
        with urllib.request.urlopen(export_url) as response, open(tmp_path, "wb") as out_f:
            out_f.write(response.read())
        return tmp_path
    except urllib.error.HTTPError as exc:
        tmp_path.unlink(missing_ok=True)
        if exc.code in (401, 403):
            raise RuntimeError(
                "Google denied anonymous access to the sheet. "
                "Download it as .xlsx from Google Sheets and pass that file path instead."
            ) from exc
        raise RuntimeError(f"Could not download Google Sheet: HTTP {exc.code}") from exc
    except urllib.error.URLError as exc:
        tmp_path.unlink(missing_ok=True)
        raise RuntimeError(f"Could not download Google Sheet: {exc.reason}") from exc


def resolve_source_workbook(source_spec):
    source_path = Path(source_spec)
    if source_path.exists():
        return source_path.resolve(), None

    downloaded = download_google_sheet_xlsx(source_spec)
    if downloaded is not None:
        return downloaded, downloaded

    raise FileNotFoundError(
        f"Source workbook not found: {source_spec}\n"
        "Pass either a local .xlsx file, a Google Sheets URL, or a Google Sheet ID."
    )


def iter_translation_rows(ws):
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        offset = row[0].value if len(row) > 0 else None
        block = row[1].value if len(row) > 1 else None
        japanese = row[4].value if len(row) > 4 else None
        english = row[6].value if len(row) > 6 else None

        if not normalize_key(japanese):
            continue

        yield RowRecord(
            sheet=ws.title,
            row_number=row_number,
            offset_key=normalize_key(offset),
            block_key=normalize_key(block),
            japanese_key=normalize_key(japanese),
            japanese_text=text_value(japanese),
            english_text=text_value(english),
        )


def build_source_index(source_path):
    wb = load_workbook(source_path, read_only=True, data_only=False)

    exact_rows = {}
    occurrence_rows = defaultdict(list)
    sheet_names = set()
    total_rows = 0
    translated_rows = 0

    try:
        for ws in wb.worksheets:
            sheet_names.add(ws.title)
            for record in iter_translation_rows(ws):
                total_rows += 1
                if is_translated(record.japanese_text, record.english_text):
                    translated_rows += 1

                exact_key = (
                    record.sheet,
                    record.offset_key,
                    record.block_key,
                    record.japanese_key,
                )
                exact_rows[exact_key] = record

                occurrence_key = (
                    record.sheet,
                    record.block_key,
                    record.japanese_key,
                )
                occurrence_rows[occurrence_key].append(record)
    finally:
        wb.close()

    return SourceIndex(
        exact_rows=exact_rows,
        occurrence_rows=dict(occurrence_rows),
        sheet_names=sheet_names,
        total_rows=total_rows,
        translated_rows=translated_rows,
    )


def backup_path_for(target_path):
    candidate = target_path.with_name(target_path.stem + ".backup" + target_path.suffix)
    if not candidate.exists():
        return candidate

    counter = 2
    while True:
        candidate = target_path.with_name(
            target_path.stem + f".backup{counter}" + target_path.suffix
        )
        if not candidate.exists():
            return candidate
        counter += 1


def merge_translations(source_index, target_path, output_path, make_backup=True, dry_run=False):
    wb = load_workbook(target_path)
    stats = Counter()
    missing_source_sheets = []
    target_occurrence_seen = defaultdict(int)

    try:
        for ws in wb.worksheets:
            if ws.title not in source_index.sheet_names:
                missing_source_sheets.append(ws.title)

            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
                japanese = row[4].value if len(row) > 4 else None
                if not normalize_key(japanese):
                    continue

                offset_key = normalize_key(row[0].value if len(row) > 0 else None)
                block_key = normalize_key(row[1].value if len(row) > 1 else None)
                japanese_key = normalize_key(japanese)

                occurrence_key = (ws.title, block_key, japanese_key)
                occurrence_index = target_occurrence_seen[occurrence_key]
                target_occurrence_seen[occurrence_key] += 1

                exact_key = (ws.title, offset_key, block_key, japanese_key)
                source_row = source_index.exact_rows.get(exact_key)
                if source_row is not None:
                    stats["exact_matches"] += 1
                else:
                    candidates = source_index.occurrence_rows.get(occurrence_key, [])
                    if occurrence_index < len(candidates):
                        source_row = candidates[occurrence_index]
                        stats["fallback_matches"] += 1
                    else:
                        stats["unmatched_rows"] += 1
                        continue

                if not is_translated(source_row.japanese_text, source_row.english_text):
                    stats["source_placeholders"] += 1
                    continue

                target_english_cell = row[6]
                current_value = text_value(target_english_cell.value)
                if current_value == source_row.english_text:
                    stats["unchanged_rows"] += 1
                    continue

                stats["updated_rows"] += 1
                if not dry_run:
                    target_english_cell.value = source_row.english_text

        if dry_run:
            return stats, missing_source_sheets, None

        backup_path = None
        if make_backup and target_path.resolve() == output_path.resolve():
            backup_path = backup_path_for(target_path)
            shutil.copy2(target_path, backup_path)

        wb.save(output_path)
        return stats, missing_source_sheets, backup_path
    finally:
        wb.close()


def parse_args(argv):
    parser = argparse.ArgumentParser(
        description="Merge English translations from a Google Sheet export into DiffRealm_Text.xlsx."
    )
    parser.add_argument(
        "source",
        help="Path to a downloaded .xlsx export, a Google Sheets URL, or a Google Sheet ID.",
    )
    parser.add_argument(
        "--target",
        default=str(DEFAULT_TARGET),
        help="Local workbook to update (default: DiffRealm_Text.xlsx).",
    )
    parser.add_argument(
        "--output",
        help="Write the merged workbook to a new path instead of updating --target in place.",
    )
    parser.add_argument(
        "--no-backup",
        action="store_true",
        help="Do not create a .backup copy when updating the target workbook in place.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Analyze matches and print a summary without writing any file.",
    )
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv or sys.argv[1:])

    target_path = Path(args.target).resolve()
    if not target_path.exists():
        print(f"Target workbook not found: {target_path}", file=sys.stderr)
        return 1

    output_path = Path(args.output).resolve() if args.output else target_path
    temp_download = None

    try:
        source_path, temp_download = resolve_source_workbook(args.source)
        print(f"Source workbook: {source_path}")
        print(f"Target workbook: {target_path}")
        if args.dry_run:
            print("Mode: dry run (no files will be written)")
        elif output_path != target_path:
            print(f"Output workbook: {output_path}")

        source_index = build_source_index(source_path)
        print(
            f"Loaded {source_index.total_rows} source rows; "
            f"{source_index.translated_rows} rows contain translated English text"
        )

        stats, missing_source_sheets, backup_path = merge_translations(
            source_index,
            target_path,
            output_path,
            make_backup=not args.no_backup,
            dry_run=args.dry_run,
        )

        print("")
        print("Merge summary:")
        print(f"  Exact row matches: {stats['exact_matches']}")
        print(f"  Fallback row matches: {stats['fallback_matches']}")
        print(f"  Rows updated: {stats['updated_rows']}")
        print(f"  Rows unchanged: {stats['unchanged_rows']}")
        print(f"  Source placeholders skipped: {stats['source_placeholders']}")
        print(f"  Target rows with no source match: {stats['unmatched_rows']}")

        if missing_source_sheets:
            print(f"  Target sheets missing from source: {len(missing_source_sheets)}")
            for sheet_name in missing_source_sheets[:10]:
                print(f"    - {sheet_name}")
            if len(missing_source_sheets) > 10:
                print(f"    ... and {len(missing_source_sheets) - 10} more")

        if args.dry_run:
            print("")
            print("Dry run complete.")
            return 0

        if backup_path is not None:
            print(f"Backup written to: {backup_path}")
        print(f"Merged workbook written to: {output_path}")
        return 0

    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    finally:
        if temp_download is not None:
            temp_download.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
