#!/usr/bin/env python
"""Build viewer data from TOS files for the Different Realm script viewer.

Usage:
    python build_script_viewer_data.py
    python build_script_viewer_data.py --no-excel   # skip Excel translation loading

Outputs script_viewer/viewer-data.js with all decoded TOS blocks.
"""

import glob
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from tos import decode_tos
from rominfo import WINDOW_WIDTH


PLAYER_NAME_JP = 'ストックマン'
PLAYER_NAME_EN = 'Stockman'
VALIDATION_PAGE_LINE_LIMIT = 3


def is_sjis_lead(b):
    return (0x81 <= b <= 0x9F) or (0xE0 <= b <= 0xEF)


def bytes_to_unicode_tagged(content):
    """Convert a block's byte content (SJIS text + ASCII tags) to a Unicode
    string with bracket-tags preserved."""
    parts = []
    pos = 0
    while pos < len(content):
        b = content[pos]
        if is_sjis_lead(b) and pos + 1 < len(content):
            sjis_pair = bytes([content[pos], content[pos + 1]])
            try:
                char = sjis_pair.decode('shift_jis')
            except Exception:
                char = f'\\x{content[pos]:02x}\\x{content[pos + 1]:02x}'
            parts.append(char)
            pos += 2
        elif b == 0x5B:  # '['
            end = pos + 1
            while end < len(content) and content[end] != 0x5D:
                end += 1
            if end < len(content):
                try:
                    tag = content[pos:end + 1].decode('ascii')
                except UnicodeDecodeError:
                    tag = '[?]'
                parts.append(tag)
                pos = end + 1
            else:
                parts.append('[')
                pos += 1
        elif 0x20 <= b <= 0x7E:
            parts.append(chr(b))
            pos += 1
        elif b in (0x0D, 0x0A):
            pos += 1
        else:
            parts.append(f'\\x{b:02x}')
            pos += 1
    return ''.join(parts)


def strip_tags_for_display(text, player_name=PLAYER_NAME_JP):
    """Strip control-code tags for display, keeping LN as newlines and
    Spaces as fullwidth spaces."""
    text = text.replace('[LN]', '\n')
    text = re.sub(r'\[Spaces(\d+)\]', lambda m: '\u3000' * int(m.group(1)), text)
    text = text.replace('[PlayerName]', player_name)
    text = re.sub(r'\[MapNameEnd\]', '', text)
    # MapName tag — just remove the tag, keep the name text that follows
    text = re.sub(r'\[MapName\]', '', text)
    text = re.sub(r'\[[^\]]+\]', '', text)
    return text


def extract_tags(raw):
    """Return a list of all tag names found in the raw text."""
    return re.findall(r'\[([^\]]+)\]', raw)


def extract_metadata(tags):
    """Extract portrait, window position, and active colors from tag list."""
    portrait = None
    window = None
    colors = set()

    for t in tags:
        if t.startswith('PortraitUp'):
            portrait = t
        elif t.startswith('PortraitDown'):
            portrait = t
        elif t.startswith('WindowUp'):
            window = 'upper'
        elif t.startswith('WindowDown'):
            window = 'lower'
        elif t.startswith('Color'):
            colors.add(t)

    return portrait, window, sorted(colors)


def summarize_preview(text, limit=36):
    cleaned = text.replace('\n', ' ')
    return cleaned if len(cleaned) <= limit else cleaned[:limit - 1] + '…'


def reconstruct_english_block(raw, entries):
    english_raw = raw
    applied = 0
    issues = []

    for entry in entries:
        japanese = entry['japanese']
        replacement = entry['english']
        if entry['suffix']:
            replacement += entry['suffix']

        if not japanese:
            continue

        if japanese not in english_raw:
            issues.append({
                'code': 'translation-segment-unmatched',
                'severity': 'error',
                'message': f'Workbook text not found in decoded block: “{summarize_preview(japanese)}”',
            })
            continue

        english_raw = english_raw.replace(japanese, replacement, 1)
        applied += 1

    return (english_raw if applied else ''), issues


def estimate_display_units(text):
    units = 0
    for char in text:
        if char == '\u3000':
            units += 2
        elif ord(char) < 0x80:
            units += 1
        else:
            units += 2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1
    return units


def parse_tagged_pages(raw, player_name=PLAYER_NAME_EN):
    pages = []
    current_page = []
    current_line_units = 0
    token_regex = re.compile(r'(\[[^\]]+\])')

    tokens = [tok for tok in token_regex.split(raw) if tok]

    for token in tokens:
        if token.startswith('[') and token.endswith(']'):
            tag = token[1:-1]
            if tag == 'LN':
                current_page.append(current_line_units)
                current_line_units = 0
            elif tag in ('Input', 'Clear'):
                current_page.append(current_line_units)
                current_line_units = 0
                pages.append(current_page)
                current_page = []
            elif tag == 'PlayerName':
                current_line_units += len(player_name)
            elif tag.startswith('Spaces'):
                try:
                    current_line_units += int(tag[6:])
                except ValueError:
                    pass
            elif tag in ('MapName', 'MapNameEnd'):
                continue
        else:
            parts = token.split('\n')
            for idx, part in enumerate(parts):
                if part:
                    current_line_units += estimate_display_units(part)
                if idx < len(parts) - 1:
                    current_page.append(current_line_units)
                    current_line_units = 0

    if current_line_units > 0 or not current_page:
        current_page.append(current_line_units)
    if current_page:
        pages.append(current_page)

    return pages or [[0]]


def build_validation(raw, english_raw, english_text, tags, portrait, base_issues):
    issues = list(base_issues)
    width_limit = None if 'MapName' in tags else WINDOW_WIDTH['PORTRAIT' if portrait else 'FULL']
    source = english_raw or english_text

    if not source:
        return {
            'widthLimit': width_limit,
            'pageLineLimit': VALIDATION_PAGE_LINE_LIMIT if width_limit is not None else None,
            'pageCount': 0,
            'rawPageCount': len(parse_tagged_pages(raw, player_name=PLAYER_NAME_JP)),
            'maxLineUnits': 0,
            'maxLinesPerPage': 0,
            'issueCount': len(issues),
            'issues': issues,
        }

    english_pages = parse_tagged_pages(source, player_name=PLAYER_NAME_EN)
    raw_pages = parse_tagged_pages(raw, player_name=PLAYER_NAME_JP)

    max_line_units = 0
    max_lines_per_page = 0

    if width_limit is not None:
        for page_index, page in enumerate(english_pages, start=1):
            max_lines_per_page = max(max_lines_per_page, len(page))
            if len(page) > VALIDATION_PAGE_LINE_LIMIT:
                issues.append({
                    'code': 'page-overflow',
                    'severity': 'warning',
                    'message': f'Page {page_index} has {len(page)} lines; dialogue windows usually fit {VALIDATION_PAGE_LINE_LIMIT}.',
                })

            for line_index, units in enumerate(page, start=1):
                max_line_units = max(max_line_units, units)
                if units > width_limit:
                    issues.append({
                        'code': 'line-overflow',
                        'severity': 'warning',
                        'message': f'Page {page_index}, line {line_index} is {units} cells wide (limit {width_limit}).',
                    })

        if len(english_pages) > len(raw_pages):
            issues.append({
                'code': 'page-count-increase',
                'severity': 'warning',
                'message': f'English expands this block to {len(english_pages)} pages (Japanese uses {len(raw_pages)}).',
            })

    return {
        'widthLimit': width_limit,
        'pageLineLimit': VALIDATION_PAGE_LINE_LIMIT if width_limit is not None else None,
        'pageCount': len(english_pages),
        'rawPageCount': len(raw_pages),
        'maxLineUnits': max_line_units,
        'maxLinesPerPage': max_lines_per_page,
        'issueCount': len(issues),
        'issues': issues,
    }


def load_translations(xlsx_path):
    """Load English translations from the workbook.  Returns a dict
    keyed by (sheet_title, block_num) → ordered translation entries."""
    translations = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            sheet = ws.title
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) < 7:
                    continue
                block_val = row[1]    # col B — block number
                jp_val = row[4]       # col E — Japanese
                en_val = row[6]       # col G — English
                suffix_val = row[8] if len(row) > 8 else None  # col I — Suffix
                if block_val is not None and en_val and jp_val and str(en_val).strip() != str(jp_val).strip():
                    key = (sheet, int(block_val))
                    if key not in translations:
                        translations[key] = []
                    translations[key].append({
                        'japanese': str(jp_val),
                        'english': str(en_val),
                        'suffix': str(suffix_val) if suffix_val else '',
                    })
        wb.close()
        print(f"  Loaded translations from {xlsx_path}: "
              f"{sum(len(v) for v in translations.values())} translated strings "
              f"across {len(translations)} blocks")
    except ImportError:
        print("  Warning: openpyxl not installed — skipping Excel translation loading")
    except FileNotFoundError:
        print(f"  Warning: {xlsx_path} not found — skipping translation loading")
    except Exception as e:
        print(f"  Warning: Could not load translations: {e}")
    return translations


def build_viewer_data(load_excel=True):
    """Decode all TOS files and build the viewer payload."""

    # Gather TOS files
    talk_files = sorted(glob.glob('original\\REALM\\TALK\\*.TOS'))
    map_files = sorted(glob.glob('original\\REALM\\MAP\\*.TOS'))
    all_files = talk_files + map_files
    all_files = [f for f in all_files
                 if '_parsed' not in f
                 and '_encoded' not in f
                 and '_roundtrip' not in f]

    print(f"Found {len(all_files)} TOS files ({len(talk_files)} TALK, {len(map_files)} MAP)")

    # Load translations
    translations = {}
    if load_excel:
        translations = load_translations('DiffRealm_Text.xlsx')

    records = []
    record_id = 0
    decode_errors = []

    for tos_file in all_files:
        name = os.path.basename(tos_file)
        stem = name.replace('.TOS', '')
        path_type = 'TALK' if '\\TALK\\' in tos_file else 'MAP'

        # Decode the TOS binary → parsed file
        parsed_file = tos_file.replace('.TOS', '_parsed.TOS')
        try:
            decode_tos(tos_file)
        except Exception as e:
            decode_errors.append(f"{name}: {e}")
            continue

        if not os.path.exists(parsed_file):
            continue

        # Read parsed blocks
        with open(parsed_file, 'rb') as f:
            lines = [l.rstrip(b'\n') for l in f.readlines()]

        for line in lines:
            if not line or line[0:1] != b'{':
                continue

            try:
                brace_end = line.index(b'}')
                block_num = int(line[1:brace_end])
            except (ValueError, IndexError):
                continue

            content = line[brace_end + 1:]
            raw = bytes_to_unicode_tagged(content)
            text = strip_tags_for_display(raw)
            tags = extract_tags(raw)
            portrait, window, colors = extract_metadata(tags)

            # Check for English translation
            english = ''
            english_raw = ''
            translation_issues = []
            key = (name, block_num)
            if key in translations:
                english_raw, translation_issues = reconstruct_english_block(raw, translations[key])
                english = strip_tags_for_display(english_raw, player_name=PLAYER_NAME_EN) if english_raw else '\n'.join(
                    entry['english'] + entry['suffix'] for entry in translations[key]
                )

            validation = build_validation(raw, english_raw, english, tags, portrait, translation_issues)

            records.append({
                'id': record_id,
                'file': name,
                'path': path_type,
                'block': block_num,
                'raw': raw,
                'text': text,
                'english': english,
                'englishRaw': english_raw,
                'portrait': portrait,
                'window': window,
                'colors': colors,
                'tags': tags,
                'validation': validation,
            })
            record_id += 1

        # Clean up parsed file
        try:
            os.remove(parsed_file)
        except OSError:
            pass

    if decode_errors:
        print(f"  Decode errors ({len(decode_errors)}):")
        for err in decode_errors:
            print(f"    {err}")

    file_names = sorted(set(r['file'] for r in records))
    translated_count = sum(1 for r in records if r['english'])
    validation_issue_count = sum(r['validation']['issueCount'] for r in records)
    validation_block_count = sum(1 for r in records if r['validation']['issueCount'])

    payload = {
        'meta': {
            'totalBlocks': len(records),
            'translatedBlocks': translated_count,
            'validationBlocks': validation_block_count,
            'validationIssues': validation_issue_count,
            'fileCount': len(file_names),
            'files': file_names,
        },
        'records': records,
    }

    return payload


def main():
    load_excel = '--no-excel' not in sys.argv

    print("Building Different Realm script viewer data...")
    data = build_viewer_data(load_excel=load_excel)

    output_dir = Path('script_viewer')
    output_dir.mkdir(exist_ok=True)

    output_path = output_dir / 'viewer-data.js'
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('window.DR_VIEWER_DATA = ')
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
        f.write(';\n')

    size_kb = output_path.stat().st_size / 1024
    meta = data['meta']
    print(f"\nDone! Wrote {output_path} ({size_kb:.1f} KB)")
    print(f"  {meta['totalBlocks']} blocks from {meta['fileCount']} files")
    print(f"  {meta['translatedBlocks']} blocks have English translations")


if __name__ == '__main__':
    main()
