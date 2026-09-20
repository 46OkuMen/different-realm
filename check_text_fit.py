"""
Report translated text that won't fit on screen, across the whole workbook.

    python check_text_fit.py              # summary per file + every fixed-field issue
    python check_text_fit.py AT01.TOS     # full detail for one file
    python check_text_fit.py --all        # full detail for everything
    python check_text_fit.py --wrap       # rewrite INFO/INFP English with word-boundary
                                          # [LN]s (edits DiffRealm_Text.xlsx)
    python check_text_fit.py --wrap-dialogue [AT01.TOS ...]
                                          # same for dialogue rows, at each row's
                                          # own window width
    python check_text_fit.py --page-breaks [AT01.TOS ...]
                                          # then add [Input][Clear]+name page breaks
                                          # wherever a page overflows 3 lines

Two kinds of check:

1. Dialogue (TALK/MAP .TOS): reuses build_script_viewer_data's per-block validation
   (line width vs WINDOW_WIDTH, lines per page, page count growth) -- the same check
   that gates reinsert.py unless --allow-validation-warnings is passed.

2. Fixed fields (menus, item/psionic info windows, data lists): the engine prints
   these into fixed-size windows and wraps character-by-character at the window
   edge, so anything too wide either breaks mid-word or collides with a value
   printed at a fixed column. FIXED_FIELDS below records each limit and where it
   was measured; the Japanese original is always the reference for what fits.
"""
import re
import sys
from collections import defaultdict

import openpyxl

from build_script_viewer_data import build_viewer_data, estimate_display_units

WORKBOOK = 'DiffRealm_Text.xlsx'
TAG = re.compile(r'\[[^\]]+\]')


def display_width(text):
    """Cells on screen for one line of tagged English text."""
    units = 0
    for tok in re.split(r'(\[[^\]]+\])', text):
        if not tok:
            continue
        if tok.startswith('[') and tok.endswith(']'):
            name = tok[1:-1]
            if name.startswith('Spaces'):
                units += 2 * int(name[6:])
            elif name == 'PlayerName':
                units += 8
            continue
        units += estimate_display_units(tok)
    return units


def wrap_breaks(text, width):
    """Simulate the engine's char-by-char wrap. Returns (line_count, mid_word_breaks)."""
    lines, breaks = 0, []
    for line in text.split('[LN]'):
        plain = TAG.sub(lambda m: ' ' * (2 * int(m.group()[7:-1])) if m.group().startswith('[Spaces') else '', line)
        lines += max(1, -(-len(plain) // width))
        for cut in range(width, len(plain), width):
            if plain[cut - 1] != ' ' and plain[cut] != ' ':
                breaks.append(plain[cut - 6:cut] + '|' + plain[cut:cut + 6])
    return lines, breaks


# sheet -> list of (selector, max_width, max_lines, note). selector is a callable on
# the workbook row's Japanese text (or None for "every row in the sheet").
CATEGORY_SLOTS = {'武 器', '　 服', '防 具', '[FW]オプション[FW]'}
FIXED_FIELDS = {
    'WORD.TOS': [
        (lambda jp: jp in CATEGORY_SLOTS, 5, 1,
         'equipment slot name, right of the item name in the Item Info header'),
        (lambda jp: jp in {'ほとんど', 'あまり', 'たまに', 'よく', '頻繁に'}, 8, 1,
         'break frequency, Item Info right panel; wraps after 8 cells (measured)'),
        (lambda jp: jp == '行動消費', 8, 1, 'Item Info right panel; the value is drawn at a fixed column after it'),
    ],
    # For max_lines == 1 rules the limit applies to every [LN]-separated line.
    'SYSTEM.TOS': [
        (lambda jp: jp == '打撃力[LN]攻撃力[LN]防御力[LN]回避力[LN]集中力[LN]抵抗力[LN]行動力', 6, 1,
         'stat labels in the 8-cell Item Info / Compare panels; "± 0" starts at col 7'),
        (lambda jp: jp in {'装備[FW]パラメータ', ' 相対評価'}, 7, 1,
         'panel title in an 8-cell window; 8 cells auto-wraps and the [LN] adds a blank line'),
        (lambda jp: jp.startswith(' 【 [Color6]装 備'), 6, 1,
         'Equipment slot labels; the equipped item name starts at col 7'),
    ],
    'ITEM.TOS': [
        (None, 15, 1, 'item name in the Item Info header (category sits at col ~18); '
                      'the Items list alone allows ~18 before the "... 1" count'),
    ],
    'INFO.TOS': [
        (None, 32, 4, 'item description window: 32 cells x 4 lines (measured in-game)'),
    ],
    'INFP.TOS': [
        (None, 32, 4, 'psionic description window: 32 cells x 4 lines (measured in-game)'),
    ],
}


def fixed_field_issues(wb):
    issues = defaultdict(list)
    for sheet, rules in FIXED_FIELDS.items():
        if sheet not in wb.sheetnames:
            continue
        for row in list(wb[sheet].iter_rows(values_only=True))[1:]:
            block, jp, en = row[1], row[4], row[6]
            if not jp or not en or en == jp:
                continue
            for selector, width, max_lines, note in rules:
                if selector is not None and not selector(jp):
                    continue
                problems = []
                widest = max(display_width(part) for part in en.split('[LN]'))
                lines, breaks = wrap_breaks(en, width)
                if max_lines == 1 and widest > width:
                    problems.append(f'{widest} cells wide (limit {width})')
                elif max_lines > 1:
                    if lines > max_lines:
                        problems.append(f'wraps to {lines} lines (window fits {max_lines})')
                    if breaks:
                        problems.append('breaks mid-word at ' + ', '.join(repr(b) for b in breaks))
                if problems:
                    issues[sheet].append((block, en, '; '.join(problems), note))
    return issues


# Sheets whose English --wrap may re-break: sheet -> (width, max_lines)
# Width is one less than the window: a line that exactly fills it auto-wraps, and the
# [LN] after it then adds a blank line.
WRAPPABLE = {'INFO.TOS': (31, 4), 'INFP.TOS': (31, 4)}


def rewrap(text, width):
    """Re-break each [LN]-separated paragraph at word boundaries so no line exceeds
    `width` cells. A breaking space becomes [LN] (both one byte encoded)."""
    # WINDOW.ASM K_RET: a line break while in Color 2 (NAME_COLOR) resets to Color 7,
    # so a break we insert inside red text must re-open it.
    def in_color2(s):
        colors = re.findall(r'\[Color(\d)\]', s)
        return bool(colors) and colors[-1] == '2'

    out = []
    for para in text.split('[LN]'):
        words = para.split(' ')
        line = ''
        for w in words:
            cand = w if not line else line + ' ' + w
            if line and display_width(cand) > width:
                out.append(line)
                line = ('[Color2]' if in_color2(''.join(out)) else '') + w
            else:
                line = cand
        out.append(line)
    return '[LN]'.join(out)


PAGE_LINES = 3          # a dialogue window shows 3 lines; a 4th scrolls the first away


def is_nametag(suffix):
    """dump_tos.py splits a speaker's name into its own row whose suffix opens the
    voice for the line that follows."""
    return bool(suffix) and 'Voice' in str(suffix)


def paginate(text, name, lines_before):
    """Insert the game's own page break -- [Input][Clear] plus the speaker's name --
    wherever a page would run past PAGE_LINES, the way the original script does it
    (AT03 block 22: name[LN]line[LN]line[Input][Clear]name[LN]...). `lines_before`
    is how many lines of this page earlier rows already used. Returns
    (text, lines_used_on_the_final_page)."""
    out, used = [], lines_before
    for i, line in enumerate(text.split('[LN]')):
        if i and used >= PAGE_LINES:
            out.append('[Input][Clear]' + (name + '[LN]' if name else '') + line)
            used = 2 if name else 1
        else:
            out.append(('[LN]' if i else '') + line)
            used += 1
        if '[Input]' in line or '[Clear]' in line:   # the author's own break
            used = 0
    return ''.join(out), used


def apply_page_breaks(sheets=None):
    """Add page breaks to dialogue that overflows its 3-line window. Without them
    the engine scrolls the first line away, so text ends up on a fresh page with no
    speaker name (and the reader may miss it). A page can span several workbook
    rows, so the line count is carried across the rows of a block and reset by any
    [Input]/[Clear], in the text or in the Suffix column. Run after
    --wrap-dialogue, since wrapping decides how many lines there are."""
    wb = openpyxl.load_workbook(WORKBOOK)
    changed = 0
    for ws in wb.worksheets:
        if ws.title in ('TagCosts',) or ws.title in FIXED_FIELDS or ws.title in WRAPPABLE:
            continue
        if sheets and ws.title not in sheets:
            continue
        block, name, used = None, None, 0
        for r in ws.iter_rows(min_row=2):
            jp, en, suffix = r[4].value, r[6].value, str(r[8].value or '')
            if not jp:
                continue
            if r[1].value != block:
                block, name, used = r[1].value, None, 0
            if is_nametag(suffix):
                name = en or jp
                used += 1
            elif en and en != jp:
                new_en, used = paginate(en, name, used)
                if new_en != en:
                    r[6].value = new_en
                    changed += 1
            else:
                used += (jp.count('[LN]') + 1)
            if '[Input]' in suffix or '[Clear]' in suffix:
                used = 0
    wb.save(WORKBOOK)
    print(f'added page breaks to {changed} rows')


def apply_wrap_dialogue(sheets=None):
    """Re-break dialogue rows at word boundaries. The engine wraps character by
    character at the window edge, so untypeset English splits words ("dangero /
    us"). Each row carries the window width the dumper measured (column C, 28 for
    portrait windows and 32 otherwise); we wrap one cell short of it, because a
    line that exactly fills the window wraps on its own and the following [LN]
    then adds a blank line. Nametag rows (their suffix opens a [Voice]) are left
    alone -- they're one short line by construction."""
    wb = openpyxl.load_workbook(WORKBOOK)
    changed = 0
    for ws in wb.worksheets:
        if ws.title in ('TagCosts',) or ws.title in FIXED_FIELDS or ws.title in WRAPPABLE:
            continue
        if sheets and ws.title not in sheets:
            continue
        for r in ws.iter_rows(min_row=2):
            jp, en, suffix = r[4].value, r[6].value, r[8].value
            if not jp or not en or en == jp:
                continue
            if suffix and 'Voice' in str(suffix):
                continue
            try:
                width = int(r[2].value)
            except (TypeError, ValueError):
                continue
            if width not in (28, 32):
                continue
            new_en = rewrap(en, width - 1)
            if new_en != en:
                r[6].value = new_en
                changed += 1
    wb.save(WORKBOOK)
    print(f'rewrapped {changed} dialogue rows')


def apply_wrap():
    wb = openpyxl.load_workbook(WORKBOOK)
    changed = 0
    for sheet, (width, max_lines) in WRAPPABLE.items():
        for r in wb[sheet].iter_rows(min_row=2):
            jp, en = r[4].value, r[6].value
            if not jp or not en or en == jp:
                continue
            new = rewrap(en, width)
            if new != en:
                r[6].value = new
                changed += 1
            lines = new.count('[LN]') + 1
            if lines > max_lines:
                print(f'{sheet} {r[1].value}: {lines} lines after wrapping (window fits {max_lines}) -- shorten: {new!r}')
    wb.save(WORKBOOK)
    print(f'rewrapped {changed} rows')


def main():
    if '--wrap' in sys.argv:
        apply_wrap()
        return
    if '--wrap-dialogue' in sys.argv:
        apply_wrap_dialogue(set(a for a in sys.argv[1:] if a.endswith('.TOS')) or None)
        return
    if '--page-breaks' in sys.argv:
        apply_page_breaks(set(a for a in sys.argv[1:] if a.endswith('.TOS')) or None)
        return

    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    show_all = '--all' in sys.argv
    only = set(args)

    data = build_viewer_data(load_excel=True)
    by_file = defaultdict(list)
    for rec in data['records']:
        if rec['file'] in WRAPPABLE:   # fixed-size info windows, checked below instead
            continue
        if rec['validation']['issueCount']:
            by_file[rec['file']].append(rec)

    print('\n=== Dialogue (per-block validation) ===')
    total = 0
    for name in sorted(by_file):
        if only and name not in only:
            continue
        recs = by_file[name]
        n = sum(r['validation']['issueCount'] for r in recs)
        total += n
        print(f'{name:12} {len(recs):4} blocks, {n:4} issues')
        if show_all or name in only:
            for r in recs:
                print(f'  block {r["block"]}: {r["english"][:70]!r}')
                for issue in r['validation']['issues']:
                    print(f'      [{issue["severity"]}] {issue["message"]}')
    print(f'total dialogue issues: {total}')

    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    print('\n=== Fixed fields (menus / info windows / data lists) ===')
    fixed = fixed_field_issues(wb)
    for sheet in sorted(fixed):
        if only and sheet not in only:
            continue
        print(f'{sheet}: {len(fixed[sheet])} issue(s) -- {fixed[sheet][0][3]}')
        for block, en, problem, _ in fixed[sheet]:
            print(f'  {str(block):>4} {en[:60]!r}: {problem}')


if __name__ == '__main__':
    main()
